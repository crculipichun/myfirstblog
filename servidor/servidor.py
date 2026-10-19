#!/usr/bin/env python3
"""Bitácora de Flota: servidor para el computador de la oficina o taller.

Los celulares envían cada registro a este computador, que lo guarda en una
planilla Excel que se actualiza sola. Con ngrok (gratis) los celulares pueden
enviar desde cualquier lugar con internet, no solo desde el Wi-Fi de la oficina.

Uso:  python servidor.py [--carpeta RUTA] [--puerto 8080] [--pin 1234]
                         [--ngrok-token TOKEN --dominio mi-flota.ngrok-free.app]
"""
import argparse
import datetime as dt
import json
import os
import re
import secrets
import socket
import sys
import threading
import time
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    print("Falta la librería openpyxl. Instálala con:  python -m pip install openpyxl")
    sys.exit(1)

try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

APP_DIR = Path(__file__).resolve().parent.parent
ARCHIVOS_APP = {"/": "index.html", "/index.html": "index.html", "/manifest.webmanifest": "manifest.webmanifest",
                "/sw.js": "sw.js", "/icon.svg": "icon.svg", "/icon-192.png": "icon-192.png", "/icon-512.png": "icon-512.png"}
COLECCIONES = ("camiones", "registros")

# Catálogos: deben coincidir con index.html
TIPOS = {"auditoria": "Auditoría", "preventiva": "Mant. preventiva", "correctiva": "Mant. correctiva",
         "falla": "Falla", "aceite": "Cambio de aceite", "odometro": "Lectura de odómetro"}
CK_ITEMS = [
    ("frenoServicio", "Frenos de servicio / balatas"), ("frenoAire", "Sistema de aire (fugas, presión)"), ("frenoEstac", "Freno de estacionamiento"),
    ("neuPresion", "Presión de inflado"), ("neuBanda", "Desgaste de banda"), ("neuPernos", "Pernos y tuercas de rueda"),
    ("luzFocos", "Focos altos y bajos"), ("luzInterm", "Intermitentes y balizas"), ("luzFreno", "Luces de freno y retroceso"),
    ("motAceite", "Nivel de aceite"), ("motRefri", "Nivel de refrigerante"), ("motFugas", "Fugas visibles"), ("motCorreas", "Correas y mangueras"),
    ("cabCinturon", "Cinturones de seguridad"), ("cabExtintor", "Extintor vigente"), ("cabKit", "Botiquín, triángulos y chaleco"),
    ("cabEspejos", "Espejos"), ("cabBocina", "Bocina y alarma de retroceso"), ("cabLimpia", "Limpiaparabrisas"),
    ("docRevision", "Revisión técnica al día"), ("docPermiso", "Permiso de circulación"), ("docSoap", "SOAP vigente"),
]
FILTROS = [("fAceite", "Filtro de aceite"), ("fAire", "Filtro de aire"), ("fComb", "Filtro de combustible"), ("fSep", "Separador de agua")]
DOCS = [("vRevision", "Revisión técnica"), ("vPermiso", "Permiso de circulación"), ("vSoap", "SOAP"), ("vExtintor", "Extintor")]


class Estado:
    def __init__(self, carpeta: Path):
        self.carpeta = carpeta
        self.json = carpeta / "datos_bitacora.json"
        self.xlsx = carpeta / "Bitacora_Flota.xlsx"
        self.lock = threading.Lock()
        self.pendiente = False
        self.ultima = None
        self.error = ""
        self.config_path = carpeta / "config_bitacora.json"
        self.config = {}
        if self.config_path.exists():
            with open(self.config_path, encoding="utf-8") as f:
                self.config = json.load(f)
        self.datos = {"camiones": {}, "registros": {}}
        if self.json.exists():
            with open(self.json, encoding="utf-8") as f:
                d = json.load(f)
            for col in COLECCIONES:
                self.datos[col] = {x["id"]: x for x in d.get(col, []) if isinstance(x, dict) and x.get("id")}
        else:
            # Primera vez: carga la flota de la empresa definida en index.html
            ahora = int(time.time() * 1000)
            self.datos["camiones"] = {c["id"]: {**c, "creado": ahora} for c in flota_inicial()}
            self.aplicar([])

    def configurar(self, **cambios):
        """Guarda la configuración; genera el PIN de administrador y el código de acceso la primera vez."""
        self.config.update({k: str(v).strip() for k, v in cambios.items() if v})
        if not self.config.get("pin_admin"):
            self.config["pin_admin"] = f"{secrets.randbelow(10**6):06d}"
        if not self.config.get("codigo_acceso"):
            letras = "ABCDEFGHJKLMNPQRSTUVWXYZ23456789"  # sin 0/O ni 1/I para que no se confundan
            self.config["codigo_acceso"] = "".join(secrets.choice(letras) for _ in range(6))
        with open(self.config_path, "w", encoding="utf-8") as f:
            json.dump(self.config, f, ensure_ascii=False, indent=1)

    @staticmethod
    def _igual(dado, esperado):
        return isinstance(dado, str) and bool(esperado) and secrets.compare_digest(dado.strip().upper(), esperado.upper())

    def es_admin(self, pin):
        return self._igual(pin, self.config.get("pin_admin", ""))

    def codigo_ok(self, codigo):
        return self._igual(codigo, self.config.get("codigo_acceso", ""))

    def lista(self):
        return {col: list(self.datos[col].values()) for col in COLECCIONES}

    def aplicar(self, cambios, admin=False):
        """Aplica los cambios. Solo el administrador puede tocar camiones; devuelve cuántos se rechazaron."""
        rechazados = 0
        for c in cambios:
            col, op = c.get("col"), c.get("op")
            if col not in COLECCIONES:
                continue
            if col == "camiones" and not admin:
                rechazados += 1
                continue
            if op == "set" and isinstance(c.get("doc"), dict) and isinstance(c["doc"].get("id"), str):
                self.datos[col][c["doc"]["id"]] = c["doc"]
            elif op == "delete" and isinstance(c.get("id"), str):
                self.datos[col].pop(c["id"], None)
        tmp = self.json.with_suffix(".tmp")
        with open(tmp, "w", encoding="utf-8") as f:
            json.dump(self.lista(), f, ensure_ascii=False, indent=1)
        os.replace(tmp, self.json)
        return rechazados

    def escribir_excel(self):
        tmp = self.carpeta / "~Bitacora_Flota_nueva.xlsx"
        try:
            construir_excel(self.lista()).save(tmp)
            os.replace(tmp, self.xlsx)
            self.pendiente, self.error, self.ultima = False, "", dt.datetime.now()
            print(f"[{self.ultima:%H:%M:%S}] Planilla actualizada: {len(self.datos['camiones'])} camiones, {len(self.datos['registros'])} registros")
        except PermissionError:
            self.pendiente = True
            self.error = "La planilla está abierta en Excel. Ciérrala y se actualizará sola."
            print(f"[{dt.datetime.now():%H:%M:%S}] {self.error}")
        except Exception as e:  # noqa: BLE001
            self.pendiente, self.error = True, f"No se pudo escribir la planilla: {e}"
            print(self.error)
        finally:
            if tmp.exists():
                try:
                    tmp.unlink()
                except OSError:
                    pass

    def info(self):
        return {"app": "bitacora-flota", "planilla": str(self.xlsx), "pendiente": self.pendiente, "error": self.error,
                "ultima": self.ultima.isoformat(timespec="seconds") if self.ultima else None}


# ---------- Cálculos (los mismos que muestra la app) ----------
def flota_inicial():
    try:
        html = (APP_DIR / "index.html").read_text(encoding="utf-8")
        m = re.search(r"/\*FLOTA_INICIAL\*/(.*?)/\*FIN_FLOTA\*/", html, re.S)
        return json.loads(m.group(1)) if m else []
    except (OSError, ValueError):
        return []


def km_de(c, regs):
    """Odómetro vigente: la lectura más alta entre la ficha y los registros del camión."""
    vals = [v for v in [c.get("km")] + [r.get("km") for r in regs if r.get("camionId") == c["id"]] if isinstance(v, (int, float))]
    return max(vals) if vals else None


def horo_de(c, regs):
    vals = [v for v in [c.get("horometro")] + [r.get("horometro") for r in regs if r.get("camionId") == c["id"]] if isinstance(v, (int, float))]
    return max(vals) if vals else None


def fecha(s):
    try:
        return dt.date.fromisoformat(s) if s else None
    except (TypeError, ValueError):
        return None


def patente(p):
    p = "".join(ch for ch in str(p or "").upper() if ch.isalnum())
    return "·".join(p[i:i + 2] for i in range(0, len(p), 2))


def orden(r):
    return (r.get("fecha") or "", r.get("km") or 0, r.get("creado") or 0)


def puntaje(ck):
    ck = ck or {}
    vals = [ck.get(k) for k, _ in CK_ITEMS if ck.get(k)]
    ok, obs, falla = vals.count("OK"), vals.count("Obs"), vals.count("Falla")
    return {"ok": ok, "obs": obs, "falla": falla, "pct": round(ok / len(vals) * 100) if vals else 0}


def titulo(r):
    t = r.get("tipo")
    if t == "auditoria":
        p = puntaje(r.get("checklist"))
        return f"Auditoría · {p['pct']}% OK" + (f" · {p['falla']} fallas" if p["falla"] else "") + (f" · {p['obs']} obs." if p["obs"] else "")
    if t == "preventiva":
        return r.get("servicio") or "Mantención preventiva"
    if t == "correctiva":
        return (r.get("sistema") or "Correctiva") + (": " + r["causa"] if r.get("causa") else "")
    if t == "falla":
        return f"Falla {(r.get('severidad') or '').lower()} · {r.get('sistema') or ''}"
    if t == "aceite":
        return f"Aceite {r.get('tipoAceite') or ''}" + (f" · {r['litros']} L" if r.get("litros") else "")
    if t == "odometro":
        return "Lectura de odómetro" + (f" · {r['horometro']} h" if r.get("horometro") else "")
    return TIPOS.get(t, t or "")


def construir_excel(d):
    camiones = {c["id"]: c for c in d["camiones"]}
    regs = sorted(d["registros"], key=orden, reverse=True)
    hoy = dt.date.today()

    def de(cid, tipo):
        return next((r for r in regs if r.get("camionId") == cid and r.get("tipo") == tipo), None)

    def abiertas(cid):
        return [r for r in regs if r.get("camionId") == cid and r.get("tipo") == "falla" and r.get("estado") != "Cerrado"]

    def aceite(c):
        u = de(c["id"], "aceite")
        if not u or u.get("proximoKm") is None:
            return u, None, "Sin registro"
        rest = u["proximoKm"] - (km_de(c, regs) or 0)
        return u, rest, "Vencido" if rest <= 0 else "Por vencer" if rest <= 1500 else "Al día"

    def doc_vencido(c):
        return any(fecha(c.get(k)) and fecha(c.get(k)) < hoy for k, _ in DOCS)

    def estado(c):
        f = abiertas(c["id"])
        if any(x.get("detenido") for x in f):
            return "Detenido"
        if f or aceite(c)[2] in ("Vencido", "Por vencer") or doc_vencido(c):
            return "Revisar"
        return "Operativo"

    def pat(r):
        return patente((camiones.get(r.get("camionId")) or {}).get("patente") or r.get("patente"))

    def base(r):
        return [fecha(r.get("fecha")), pat(r), r.get("km")]

    sn = lambda b: "Sí" if b else "No"  # noqa: E731
    wb = Workbook()
    wb.remove(wb.active)

    def hoja(nombre, cab, filas):
        ws = wb.create_sheet(nombre)
        ws.append(cab)
        for f in filas:
            ws.append(["" if v is None else v for v in f])
        amarillo = PatternFill("solid", fgColor="F2A900")
        for cell in ws[1]:
            cell.font, cell.fill = Font(bold=True, color="1D1500"), amarillo
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.row_dimensions[1].height = 32
        ws.freeze_panes = "A2"
        if filas:
            ws.auto_filter.ref = ws.dimensions
        for i, h in enumerate(cab, 1):
            largo = max([len(str(h))] + [len(str(f[i - 1])) for f in filas if f[i - 1] is not None][:500])
            ws.column_dimensions[get_column_letter(i)].width = min(48, max(10, largo + 2))
            for (cell,) in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                if isinstance(cell.value, dt.date):
                    cell.number_format = "DD-MM-YYYY"
                elif isinstance(cell.value, (int, float)) and ("Km" in h or "km" in h or "Costo" in h):
                    cell.number_format = '"$"#,##0' if "Costo" in h else "#,##0"

    filas = []
    def interno(c):
        try:
            return (0, int(c.get("interno")), patente(c.get("patente")))
        except (TypeError, ValueError):
            return (1, 0, patente(c.get("patente")))

    for c in sorted(camiones.values(), key=interno):
        u, rest, est_ac = aceite(c)
        au = de(c["id"], "auditoria")
        filas.append([patente(c.get("patente")), c.get("interno"), c.get("tipoCamion"), c.get("marca"), c.get("modelo"), c.get("anio"),
                      c.get("color"), c.get("vin"), c.get("motor"), km_de(c, regs), horo_de(c, regs), c.get("conductor"), c.get("faena"),
                      c.get("intervalo"), fecha(u.get("fecha")) if u else None, u.get("km") if u else None, u.get("proximoKm") if u else None,
                      rest, est_ac, len(abiertas(c["id"])), fecha(au.get("fecha")) if au else None, puntaje(au.get("checklist"))["pct"] if au else None,
                      *[fecha(c.get(k)) for k, _ in DOCS], c.get("estanque"), c.get("neumaticos"), estado(c), c.get("notas")])
    hoja("Flota", ["Patente", "N° interno", "Tipo", "Marca", "Modelo", "Año", "Color", "VIN / N° chasis", "N° motor", "Odómetro (km)",
                   "Horómetro (h)", "Conductor", "Faena / centro de costo", "Intervalo aceite (km)", "Último cambio aceite", "Km último cambio",
                   "Próximo cambio (km)", "Km restantes", "Estado aceite", "Fallas abiertas", "Última auditoría", "% OK última auditoría",
                   *[f"Vence {l}" for _, l in DOCS], "Estanque (L)", "Neumáticos", "Estado", "Notas"], filas)

    hoja("Auditorías", ["Fecha", "Patente", "Odómetro (km)", "Auditor", "% OK", "OK", "Obs", "Fallas", *[l for _, l in CK_ITEMS], "Observaciones", "Estado"],
         [[*base(r), r.get("responsable"), *(lambda p: [p["pct"], p["ok"], p["obs"], p["falla"]])(puntaje(r.get("checklist"))),
           *[(r.get("checklist") or {}).get(k, "") for k, _ in CK_ITEMS], r.get("descripcion"), r.get("estado")]
          for r in regs if r.get("tipo") == "auditoria"])
    hoja("Preventivas", ["Fecha", "Patente", "Odómetro (km)", "Responsable", "Servicio", "Repuestos", "Próximo (km)", "Próxima fecha", "Costo CLP", "Estado", "Observaciones"],
         [[*base(r), r.get("responsable"), r.get("servicio"), r.get("repuestos"), r.get("proximoKm"), fecha(r.get("proximaFecha")), r.get("costo"), r.get("estado"), r.get("descripcion")]
          for r in regs if r.get("tipo") == "preventiva"])
    hoja("Correctivas", ["Fecha", "Patente", "Odómetro (km)", "Responsable", "Sistema", "Causa", "Repuestos", "Horas taller", "Taller", "Costo CLP", "Estado", "Fecha cierre", "Observaciones"],
         [[*base(r), r.get("responsable"), r.get("sistema"), r.get("causa"), r.get("repuestos"), r.get("horas"), r.get("taller"), r.get("costo"), r.get("estado"), fecha(r.get("cerrado")), r.get("descripcion")]
          for r in regs if r.get("tipo") == "correctiva"])
    hoja("Fallas", ["Fecha", "Patente", "Odómetro (km)", "Reportado por", "Sistema", "Severidad", "Camión detenido", "Estado", "Fecha cierre", "Descripción"],
         [[*base(r), r.get("responsable"), r.get("sistema"), r.get("severidad"), sn(r.get("detenido")), r.get("estado"), fecha(r.get("cerrado")), r.get("descripcion")]
          for r in regs if r.get("tipo") == "falla"])
    hoja("Cambios de aceite", ["Fecha", "Patente", "Odómetro (km)", "Responsable", "Aceite", "Litros", *[l for _, l in FILTROS], "Próximo cambio (km)", "Costo CLP", "Observaciones"],
         [[*base(r), r.get("responsable"), r.get("tipoAceite"), r.get("litros"), *[sn(r.get(k)) for k, _ in FILTROS], r.get("proximoKm"), r.get("costo"), r.get("descripcion")]
          for r in regs if r.get("tipo") == "aceite"])
    hoja("Odómetro", ["Fecha", "Patente", "Odómetro (km)", "Horómetro (h)", "Responsable", "Observaciones"],
         [[*base(r), r.get("horometro"), r.get("responsable"), r.get("descripcion")] for r in regs if r.get("tipo") == "odometro"])
    hoja("Todos los registros", ["Fecha", "Tipo", "Patente", "Odómetro (km)", "Responsable", "Detalle", "Estado", "Costo CLP", "Observaciones"],
         [[fecha(r.get("fecha")), TIPOS.get(r.get("tipo"), r.get("tipo")), pat(r), r.get("km"), r.get("responsable"), titulo(r), r.get("estado"), r.get("costo"), r.get("descripcion")]
          for r in regs])
    return wb


# ---------- Servidor web ----------
def hacer_handler(estado: Estado):
    class Handler(SimpleHTTPRequestHandler):
        def __init__(self, *a, **k):
            super().__init__(*a, directory=str(APP_DIR), **k)

        def log_message(self, *a):
            pass

        def responder(self, codigo, obj):
            cuerpo = json.dumps(obj, ensure_ascii=False).encode("utf-8")
            self.send_response(codigo)
            self.send_header("Content-Type", "application/json; charset=utf-8")
            self.send_header("Cache-Control", "no-store")
            self.send_header("Content-Length", str(len(cuerpo)))
            self.end_headers()
            self.wfile.write(cuerpo)

        def autorizado(self):
            """Las consultas de datos exigen el código de acceso de la empresa (cabecera X-Codigo)."""
            if estado.codigo_ok(self.headers.get("X-Codigo")):
                return True
            time.sleep(1)  # frena a quien prueba códigos al azar
            self.responder(401, {"error": "codigo", "mensaje": "Código de acceso incorrecto"})
            return False

        def do_GET(self):
            ruta = self.path.split("?")[0]
            if ruta == "/api/estado":
                # Sin código solo confirma que aquí está la app; con código entrega el estado de la planilla
                if estado.codigo_ok(self.headers.get("X-Codigo")):
                    return self.responder(200, {**estado.info(), "codigoOk": True})
                return self.responder(200, {"app": "bitacora-flota", "codigoOk": False})
            if ruta == "/api/datos":
                if not self.autorizado():
                    return
                with estado.lock:
                    return self.responder(200, {"datos": estado.lista(), "planilla": estado.info()})
            if ruta not in ARCHIVOS_APP:
                return self.send_error(404)
            self.path = "/" + ARCHIVOS_APP[ruta]
            return super().do_GET()

        def end_headers(self):
            if not self.path.startswith("/api/"):
                self.send_header("Cache-Control", "no-cache")
            # La app Android carga la página desde el teléfono y llama a este computador
            self.send_header("Access-Control-Allow-Origin", "*")
            super().end_headers()

        def do_OPTIONS(self):
            self.send_response(204)
            self.send_header("Access-Control-Allow-Methods", "GET, POST, OPTIONS")
            self.send_header("Access-Control-Allow-Headers", "Content-Type, X-Codigo, ngrok-skip-browser-warning")
            self.send_header("Access-Control-Allow-Private-Network", "true")
            self.send_header("Access-Control-Max-Age", "86400")
            self.end_headers()

        def leer_json(self):
            largo = int(self.headers.get("Content-Length") or 0)
            if largo > 20 * 1024 * 1024:
                raise ValueError("demasiado grande")
            d = json.loads(self.rfile.read(largo) or b"{}")
            if not isinstance(d, dict):
                raise ValueError
            return d

        def do_POST(self):
            ruta = self.path.split("?")[0]
            if ruta not in ("/api/sync", "/api/admin"):
                return self.send_error(404)
            try:
                cuerpo = self.leer_json()
                cambios = cuerpo.get("cambios") or []
                if not isinstance(cambios, list):
                    raise ValueError
            except ValueError:
                return self.responder(400, {"error": "formato inválido"})
            if not self.autorizado():
                return
            admin = estado.es_admin(cuerpo.get("pin"))
            if ruta == "/api/admin":
                if not admin:
                    time.sleep(1)  # frena a quien prueba PIN al azar
                return self.responder(200, {"ok": admin})
            with estado.lock:
                rechazados = 0
                if cambios:
                    rechazados = estado.aplicar(cambios, admin)
                    estado.escribir_excel()
                return self.responder(200, {"datos": estado.lista(), "planilla": estado.info(), "admin": admin, "rechazados": rechazados})

    return Handler


def ip_local():
    s = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
    try:
        s.connect(("10.255.255.255", 1))
        return s.getsockname()[0]
    except OSError:
        return "127.0.0.1"
    finally:
        s.close()


def abrir_ngrok(estado: Estado, puerto):
    """Publica el servidor en internet con la dirección fija de ngrok. Devuelve la URL o None."""
    token, dominio = estado.config.get("ngrok_token"), estado.config.get("ngrok_dominio")
    if not token:
        return None
    try:
        from pyngrok import conf, ngrok
    except ImportError:
        print(" Falta pyngrok para conectarse por internet. Instálalo con:  python -m pip install pyngrok")
        return None
    try:
        conf.get_default().auth_token = token
        opciones = {"domain": dominio} if dominio else {}
        return ngrok.connect(puerto, "http", **opciones).public_url
    except Exception as e:  # noqa: BLE001
        print(f" No se pudo abrir la conexión por internet (ngrok): {e}")
        return None


def reintentar(estado: Estado):
    while True:
        time.sleep(15)
        if estado.pendiente:
            with estado.lock:
                estado.escribir_excel()


def main():
    ap = argparse.ArgumentParser(description="Servidor de Bitácora de Flota")
    ap.add_argument("--carpeta", default=str(Path.home() / "Documents" / "Bitacora Flota"), help="Carpeta donde se guarda la planilla Excel")
    ap.add_argument("--puerto", type=int, default=8080)
    ap.add_argument("--pin", help="Cambia el PIN de administrador")
    ap.add_argument("--codigo", help="Cambia el código de acceso que usan los celulares")
    ap.add_argument("--ngrok-token", help="Authtoken de tu cuenta ngrok (se guarda, basta darlo una vez)")
    ap.add_argument("--dominio", help="Dominio fijo gratuito de ngrok, por ejemplo mi-flota.ngrok-free.app")
    args = ap.parse_args()
    carpeta = Path(args.carpeta).expanduser()
    carpeta.mkdir(parents=True, exist_ok=True)
    estado = Estado(carpeta)
    estado.configurar(pin_admin=args.pin, codigo_acceso=args.codigo, ngrok_token=args.ngrok_token, ngrok_dominio=args.dominio)
    cfg = estado.config
    with estado.lock:
        estado.escribir_excel()
    threading.Thread(target=reintentar, args=(estado,), daemon=True).start()
    srv = ThreadingHTTPServer(("0.0.0.0", args.puerto), hacer_handler(estado))
    publica = abrir_ngrok(estado, args.puerto)
    print("=" * 66)
    print(" BITÁCORA DE FLOTA · servidor del computador")
    print(f" Planilla Excel:  {estado.xlsx}")
    if publica:
        print(f" Dirección por internet (cualquier celular, cualquier lugar):")
        print(f"     {publica}")
        print(f" Link para enviar a los celulares (incluye el código):")
        print(f"     {publica}/?c={cfg['codigo_acceso']}")
    else:
        print(" Solo red local. Para recibir datos desde cualquier lugar, configura ngrok (ver README).")
    print(f" En la misma red Wi-Fi:  http://{ip_local()}:{args.puerto}/?c={cfg['codigo_acceso']}")
    print(f" Código de acceso para los celulares:  {cfg['codigo_acceso']}")
    print(f" PIN de administrador (agrega y edita camiones):  {cfg['pin_admin']}")
    print(" Deja esta ventana abierta. Para detener: Ctrl+C")
    print("=" * 66)
    try:
        srv.serve_forever()
    except KeyboardInterrupt:
        print("\nServidor detenido.")


if __name__ == "__main__":
    main()
